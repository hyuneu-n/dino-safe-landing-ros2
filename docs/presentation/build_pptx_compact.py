#!/usr/bin/env python3
"""Build the compact PowerPoint from finalized HTML renders and embed notes."""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


HERE = Path(__file__).resolve().parent
PNG_DIR = HERE / "render_compact_png"
NOTES_FILE = HERE / "speaker_notes_compact.md"
OUTPUT = HERE / "safe_landing_review_compact.pptx"


def read_notes() -> list[str]:
    source = NOTES_FILE.read_text(encoding="utf-8")
    headings = list(re.finditer(r"^##\s+\d+\.\s+.+$", source, re.MULTILINE))
    notes: list[str] = []
    for index, match in enumerate(headings):
        start = match.end()
        end = headings[index + 1].start() if index + 1 < len(headings) else len(source)
        notes.append(source[start:end].strip())
    if len(notes) != 10:
        raise RuntimeError(f"Expected 10 note sections, found {len(notes)}")
    return notes


def build() -> Path:
    images = sorted(PNG_DIR.glob("slide_??.png"))
    if len(images) != 10:
        raise RuntimeError(f"Expected 10 rendered slides in {PNG_DIR}, found {len(images)}")
    notes = read_notes()

    deck = Presentation()
    deck.slide_width = Inches(16)
    deck.slide_height = Inches(9)
    blank = deck.slide_layouts[6]

    for image, note in zip(images, notes):
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image), 0, 0, width=deck.slide_width, height=deck.slide_height
        )
        slide.notes_slide.notes_text_frame.text = note

    deck.core_properties.title = "종합설계 — RGB-D 기반 배송 드론 안전 착륙 시스템"
    deck.core_properties.subject = "종합설계 발표"
    deck.core_properties.author = "서현은"
    deck.core_properties.comments = (
        "HTML 확정 화면을 이미지로 배치한 레이아웃 고정형 PPTX. "
        "발표 대본은 각 슬라이드의 발표자 노트에 포함됨."
    )
    deck.save(OUTPUT)
    print(f"Created {OUTPUT} with {len(deck.slides)} slides and embedded notes")
    return OUTPUT


if __name__ == "__main__":
    build()
