#!/usr/bin/env python3
"""Package the rendered V2 HTML slides into a layout-stable PowerPoint."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


HERE = Path(__file__).resolve().parent
PNG_DIR = HERE / "render_v2_png"
OUTPUT = HERE / "safe_landing_first_review_v2.pptx"
PDF_OUTPUT = HERE / "safe_landing_first_review_v2.pdf"


def build() -> Path:
    images = sorted(PNG_DIR.glob("slide_*.png"))
    if len(images) != 23:
        raise RuntimeError(f"Expected 23 rendered slides in {PNG_DIR}, found {len(images)}")

    deck = Presentation()
    deck.slide_width = Inches(15)
    deck.slide_height = Inches(9)
    blank = deck.slide_layouts[6]

    for image in images:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image),
            0,
            0,
            width=deck.slide_width,
            height=deck.slide_height,
        )

    # Remove the starter slide created by python-pptx, if present.
    if len(deck.slides) == len(images) + 1:
        slide_id = deck.slides._sldIdLst[0]
        relationship_id = slide_id.rId
        deck.part.drop_rel(relationship_id)
        del deck.slides._sldIdLst[0]

    deck.core_properties.title = "RGB-D 기반 배송 드론 안전 착륙 시스템"
    deck.core_properties.subject = "종합설계 1차 발표 V2"
    deck.core_properties.comments = (
        "화면 고정형 PPTX입니다. 내용 수정은 safe_landing_first_review_v2.html과 "
        "build_html_deck_v2.py에서 진행하세요."
    )
    deck.save(OUTPUT)
    pdf_pages = [Image.open(path).convert("RGB") for path in images]
    pdf_pages[0].save(
        PDF_OUTPUT,
        save_all=True,
        append_images=pdf_pages[1:],
        resolution=120,
    )
    for page in pdf_pages:
        page.close()
    print(f"Created {OUTPUT} with {len(deck.slides)} slides")
    print(f"Created {PDF_OUTPUT} with {len(images)} pages")
    return OUTPUT


if __name__ == "__main__":
    build()
