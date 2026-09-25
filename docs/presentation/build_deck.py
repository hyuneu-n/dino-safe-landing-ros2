#!/usr/bin/env python3
"""Editable PPTX + raster PDF/PNG previews from one layout specification.
Requires python-pptx, Pillow; Korean font defaults to Windows Malgun Gothic.
Run: PYTHONPATH=/tmp/safe_landing_presentation_deps python3 docs/presentation/build_deck.py
"""
import os
from pathlib import Path
import sys
import json
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
FONT=Path(os.environ.get('PRESENTATION_FONT','/mnt/c/Windows/Fonts/malgun.ttf'))
BOLD=Path(os.environ.get('PRESENTATION_BOLD_FONT','/mnt/c/Windows/Fonts/malgunbd.ttf'))
if not FONT.exists():raise SystemExit('Set PRESENTATION_FONT and PRESENTATION_BOLD_FONT to Korean TTF files')
W,H=1280,720
BG='#101B29';CARD='#1A2B3D';WHITE='#F2F1EB';MUTED='#AEBFCC';GOLD='#F0BA58';TEAL='#6BD0BD'
prs=Presentation();prs.slide_width=Inches(W/96);prs.slide_height=Inches(H/96)
prs.core_properties.title='RGB-D 기반 배송 드론 안전 착륙지 판단 — 종합설계 1차 발표'
prs.core_properties.subject='프로젝트 소개와 현재 구현·검증 범위'
prs.core_properties.author='safe_landing'
slides=[];notes=[];alltext=[]

def col(c):return RGBColor.from_string(c.lstrip('#'))
def rect(x,y,w,h,c):
    d.rectangle((x,y,x+w,y+h),fill=c)
    shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x/96),Inches(y/96),Inches(w/96),Inches(h/96))
    shape.fill.solid();shape.fill.fore_color.rgb=col(c);shape.line.fill.background()
    return shape

def text(value,x,y,w,size=28,color=WHITE,bold=False):
    font=ImageFont.truetype(str(BOLD if bold else FONT),size)
    lines=[]
    for para in value.split('\n'):
        line=''
        for ch in para:
            if line and d.textlength(line+ch,font=font)>w:
                lines.append(line);line=ch
            else:line+=ch
        lines.append(line)
    value='\n'.join(lines);height=len(lines)*size*1.38+12
    assert y+height<=720,(len(slides)+1,value,y,height)
    box=s.shapes.add_textbox(Inches(x/96),Inches(y/96),Inches(w/96),Inches(height/96))
    tf=box.text_frame;tf.clear();tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name='맑은 고딕';p.font.size=Pt(size*.75);p.font.bold=bold;p.font.color.rgb=col(color)
        p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(size*1.38*.75)
        d.text((x,y+i*size*1.38),line,font=font,fill=color)
    alltext.append(value)
    return height

def photo(name,x,y,w,h):
    path=(OUT/name[1:]) if name.startswith('@') else ROOT/'docs/images'/name
    cropped=ImageOps.fit(Image.open(path).convert('RGB'),(int(w),int(h)))
    im.paste(cropped,(int(x),int(y)))
    # Preserve editable picture placement using a generated crop asset.
    cache=OUT/'_assets';cache.mkdir(exist_ok=True)
    target=cache/f'{len(slides):02}_{len(s.shapes):02}.jpg';cropped.save(target,quality=94)
    s.shapes.add_picture(str(target),Inches(x/96),Inches(y/96),width=Inches(w/96),height=Inches(h/96))

def start(title,kicker='SAFE LANDING / 종합설계',subtitle=None):
    global s,im,d
    s=prs.slides.add_slide(prs.slide_layouts[6]);im=Image.new('RGB',(W,H),BG);d=ImageDraw.Draw(im)
    s.background.fill.solid();s.background.fill.fore_color.rgb=col(BG)
    rect(48,42,42,4,GOLD);text(kicker,105,29,1100,17,MUTED)
    text(title,48,77,1184,43,WHITE,True)
    if subtitle:text(subtitle,50,142,1160,22,MUTED)

def end(note):
    text('서경대학교 · 종합설계 | 구현 기준 2026.09.15',48,677,1000,14,MUTED)
    text(f'{len(slides)+1:02}',1175,675,70,17,GOLD)
    s.notes_slide.notes_text_frame.text=note
    slides.append(im.copy());notes.append(note)

def card(x,y,w,h,heading,body):
    rect(x,y,w,h,CARD);text(heading,x+22,y+20,w-44,29,GOLD,True)
    text(body,x+22,y+76,w-44,25,WHITE)

start('좌표에 장애물이 있다면, 어디에 내려야 할까?',subtitle='RGB-D 기반 배송 드론 안전 착륙지 판단')
photo('metropolis/campus_aerial.jpg',48,204,742,413)
text('safe_landing',837,218,390,38,GOLD,True)
text('프로젝트 소개\n현재 구현 결과\n검증 범위와 후속 논의',837,296,380,29)
text('종합설계 1차 발표\n발표자: [이름 입력]',837,510,380,22,MUTED)
end('이 프로젝트는 배송 드론이 목적지 근처에서 안전하게 내릴 자리를 판단하는 시스템입니다. 목적지 좌표가 주어져도 그 자리가 비어 있다는 보장은 없습니다. 저는 RGB와 깊이 영상을 결합해 현재 장면의 장애물과 빈 공간을 판단하는 기능을 만들고 있습니다. 오늘은 프로젝트 목적부터 현재 구현과 검증 범위까지 소개하겠습니다.')

start('목적지 좌표와 실제 착륙점은 다를 수 있다',subtitle='같은 장소라도 사람·차량·화물에 따라 착륙 가능 공간이 달라진다')
photo('metropolis/freight_yard_down.png',48,205,540,365)
card(625,205,605,165,'요청: 이 위치로 배송','실제 현장: 지정 위치에 화물 또는 사람')
card(625,390,605,180,'필요한 판단','주변 빈 공간을 선택하거나\n안전한 자리가 없으면 기다리기')
text('핵심 문제: 좌표 도달을 넘어, 현재 장면에 맞는 착륙 공간 선택',50,608,1180,26,TEAL)
end('화면은 실제 Gazebo 하강 카메라로 본 물류 야드입니다. 배송 좌표와 현재 안전한 착륙점은 다를 수 있습니다. 따라서 좌표까지 이동하는 항법과, 도착한 현장에서 내릴 자리를 판단하는 문제를 구분했습니다. 특히 사람이 접근하거나 공간이 점유되면 다른 곳을 선택하거나 기다리는 동작이 최종 목표입니다.')

start('프로젝트 목표: 장면을 보고 착륙점을 결정',subtitle='목표 시연 흐름 — 전체 통합 동작은 아직 검증 전')
for i,(h,b) in enumerate([('01 배송지 선택','관람객이 장소 선택'),('02 이동·스캔','전방과 하강 카메라'),('03 공간 판단','장애물·여유 공간'),('04 착륙·대기','상황 변화 시 재평가')]):
    card(48+i*300,226,280,230,h,b)
text('현재 연구 중심',50,513,400,28,GOLD,True)
text('RGB-D 기반 장애물 판단과 안전 착륙점 선택',50,559,1140,34,WHITE,True)
end('최종적으로는 관람객이 배송지를 고르면 드론이 이동하고 현장을 스캔해서 착륙점을 정하도록 만들려고 합니다. 관람객은 사람 이동 같은 변화도 줄 수 있습니다. 다만 이 흐름이 현재 모두 완성된 것은 아닙니다. 지금 연구의 중심은 RGB-D로 장애물을 판단하고 안전한 착륙점을 선택하는 부분입니다.')

start('두 종류의 시각 정보를 함께 사용한다',subtitle='외관의 차이와 공간의 차이를 결합하는 접근')
card(48,210,563,265,'RGB + DINOv2','패치 특징을 군집화해\n표면의 시각적 차이를 구분\n\n사람·차량 이름을 분류하는 방식은 아님')
card(635,210,595,265,'Depth · 깊이 영상','카메라와 표면 사이 거리로\n높이 차이와 장애물 후보를 판단\n\n후보 주변의 공간 여유를 고려')
rect(48,518,1182,103,'#224A4B');text('결합 → 후보 영역 생성 → 점수화 → 착륙점 선택',76,547,1110,31,WHITE,True)
end('RGB에서는 DINOv2 특징을 추출하고 군집화해 표면의 시각적 차이를 구분합니다. 깊이 영상에서는 거리와 높이 차이를 이용합니다. 두 정보를 결합해 후보를 만들고 점수를 계산합니다. 여기서 DINO 군집은 사람이나 자동차의 이름을 붙이는 객체 검출기가 아니라는 점을 구분해서 설명하겠습니다.')

start('현재 시스템은 인식·미션·시각화로 구성',subtitle='ROS 2 Humble + Gazebo Classic 11')
card(48,209,350,185,'Gazebo','RGB · depth · odom\n도시 환경과 드론 모델')
card(465,209,350,185,'착륙 판단 노드','DINOv2 + depth\n후보와 선택점 발행')
card(882,209,348,185,'미션 상태머신','경유점 이동 · 스캔\n선택점으로 하강')
text('→',411,270,50,37,GOLD);text('→',830,270,50,37,GOLD)
rect(48,433,1182,74,CARD);text('RViz: 카메라 · 안전 영역 · 후보 · 선택점 · 미션 상태',73,453,1110,28,TEAL)
text('현재 위치 입력: Gazebo odom  /  이동: 위치를 직접 지정하는 모델',50,546,1160,26)
text('완전한 비전 전용 위치 추정·실기체 비행 제어는 별도 범위',50,594,1160,24,MUTED)
end('Gazebo가 RGB, 깊이 영상과 위치를 제공합니다. 착륙 판단 노드가 후보를 만들고 선택점을 보내면 미션 상태머신이 이동과 하강을 수행합니다. RViz로 중간 결과를 볼 수 있습니다. 현재 자기 위치는 시뮬레이터 odom을 사용하고, 이동도 위치 지정 모델입니다. 따라서 완전한 비전 전용 위치 추정이나 실제 드론 비행 제어까지 완성했다고 주장하지 않습니다.')

start('실험 기록으로 착륙점 선택을 확인',subtitle='기존 평가 seed 0 · 저장된 RGB와 선택 좌표를 함께 표시')
photo('@source_images/seed0_rgb.png',48,210,565,365);photo('@source_images/seed0_pick.png',650,210,580,365)
text('입력: 저장된 하강 RGB 프레임',50,593,555,27,GOLD)
text('출력: 융합 방식의 기록된 선택점',651,593,550,27,TEAL)
end('왼쪽은 이전 실험 seed0에 저장된 실제 하강 RGB이고, 오른쪽은 같은 프레임에 실험 JSON의 융합 방식 선택 좌표를 노란 표식으로 표시한 것입니다. 새로 알고리즘을 실행하거나 안전 마스크를 만들어 넣은 것이 아니라 저장된 결과를 시각화했습니다. 주변 장애물과 선택점의 관계를 볼 수 있습니다. 이 기록 역시 현재 도시의 비행 성공을 뜻하지 않습니다.')

start('다양한 환경을 만든 이유는 판단 조건의 변화',subtitle='환경 제작 마일스톤 완료 · 사용자가 현재 맵 구성에 만족한 상태')
for i,(name,label) in enumerate([('downtown.png','고층 도심 / 좁은 공간'),('industrial.png','물류 지구 / 적재물'),('interchange.png','고가도로 / 높이 차이'),('mountain_town.png','산동네 / 경사 지형')]):
    x=48+(i%2)*605;y=202+(i//2)*219
    photo('metropolis/'+name,x,y,578,171);text(label,x,y+178,570,23,GOLD)
end('맵은 보기 좋은 배경뿐 아니라 조건이 바뀌는 시험 환경을 만들기 위해 구성했습니다. 도심은 높은 구조물과 좁은 공간, 물류 야드는 화물과 차량, 교량과 산동네는 높이와 경사가 다릅니다. 현재 환경 제작은 마무리했고, 다음에는 이런 조건에서 알고리즘을 안정적으로 검증하는 단계로 넘어가려 합니다.')

start('도시 안에 배송 출발지와 알아볼 수 있는 목적지',subtitle='약 2.5 × 1.9km  ·  건물 148동  ·  목적지 22곳  ·  지정 위치 점유 상황 4곳')
photo('metropolis/logistics.png',48,212,565,353);photo('metropolis/campus.png',650,212,580,353)
text('SkyDrop 물류창고: 드론 출발점',50,589,570,26,GOLD)
text('사진 기반 서경대학교 캠퍼스',650,589,580,26,TEAL)
end('현재 seed0 기준으로 건물 148동과 목적지 22곳이 있습니다. 드론 출발점은 물류창고에 연결했습니다. 캠퍼스는 실제 학교 사진의 유리 타워, 아치 창문, 노란 광장과 계단 정원을 반영했습니다. 실측 복원은 아니지만 관람객이 목적지를 쉽게 이해하도록 구성했습니다. 맥도날드와 써브웨이 같은 상점도 있습니다.')

start('60초 환경 투어',subtitle='실제 Gazebo RGB 렌더링 · 환경 소개용 카메라 경로')
poster=ROOT/'docs/images/metropolis/tour_poster.jpg'
photo('metropolis/tour_poster.jpg',240,185,800,450)
video=ROOT/'artifacts/metropolis_tour/metropolis_flythrough.mp4'
s.shapes.add_movie(str(video),Inches(240/96),Inches(185/96),Inches(800/96),Inches(450/96),poster_frame_image=str(poster),mime_type='video/mp4')
text('슬라이드 쇼에서 영상을 클릭해 재생  ·  자율 회피 비행의 검증 영상은 아님',135,638,1020,18,GOLD)
end('지금부터 60초 동안 전체 환경을 보여드리겠습니다. 물류창고에서 출발해 도심, 공장, 마을, 강과 고가도로, 캠퍼스, 산동네를 훑습니다. 이 영상은 전용 촬영 카메라의 경로를 지정해 만든 환경 소개 영상입니다. 비전 알고리즘이 이 경로를 자율적으로 비행한 결과와는 구분합니다. [영상 60초 재생]')

start('현재까지 구현한 것과 아직 남은 것',subtitle='완성률 대신 실제 동작과 검증 범위로 설명')
rows=[('도시·센서 환경','구축 및 배치·RGB/depth 검사 완료',TEAL),('착륙점 탐지','DINOv2 + depth, 후보/선택점 시각화 구현',TEAL),('미션·전방 회피','기본 상태머신과 반응형 좌우 회피 구현',GOLD),('도시 전체 배송','종단 간 비행 성공률 미측정',GOLD),('대시보드·반복 배송','시연 구상 단계, 아직 미구현',MUTED)]
for i,(a,b,c) in enumerate(rows):
 y=205+i*78;rect(48,y,1182,66,CARD);text(a,69,y+16,300,25,c,True);text(b,390,y+16,810,25)
end('환경과 센서 검사는 완료했고, 착륙점 탐지와 시각화도 구현되어 있습니다. 미션 상태머신과 전방 회피는 기본 동작이 있습니다. 그러나 새로운 도시의 모든 목적지에 대해 비행과 착륙 성공률을 측정한 것은 아닙니다. 대시보드와 반복 배송도 아직 구현 전입니다. 따라서 현재는 환경과 프로토타입을 갖추고 통합 신뢰성을 높일 단계입니다.')

start('이전 실험: 24개 장면의 착륙점 선택 비교',subtitle='정지 프레임 평가 · 새 도시의 배송 성공률이 아님')
for i,(name,n,dist,c) in enumerate([('DINOv2 + depth',22,'3.69m',TEAL),('Depth-only 재구현',17,'2.22m',GOLD),('OpenLander RGB-only',11,'2.63m',MUTED)]):
 y=214+i*119;text(name,50,y,360,27,c,True);rect(445,y+4,620,32,CARD);rect(445,y+4,620*n/24,32,c)
 text(f'{n}/24',1100,y,125,28,c,True);text('평균 MOD '+dist,445,y+48,620,22,MUTED)
text('성공 기준: 선택점의 반경 1.0m 안에 GT 장애물 없음',50,590,1180,24)
text('원자료 24행 재확인 · 후보 제시율: 24/24, 24/24, 20/24',50,634,1180,18,MUTED)
end('이전 단일 회랑 월드에서 24개 정지 장면을 비교했습니다. 선택점 주변 반경 1미터 안에 정답 장애물이 없는 경우를 성공으로 보았습니다. 융합 방식은 22개, depth-only는 17개, RGB-only 비교 방식은 11개에서 성공했습니다. 이는 착륙점 선택 평가이지 실제 하강이나 비행 전체 성공률이 아닙니다. OpenLander는 네 장면에서 후보를 제시하지 않았으며, 평균 장애물 거리는 후보가 나온 경우에 대해 계산했습니다.')

start('다음 단계는 통합 신뢰성 검증',subtitle='기술 선택과 우선순위는 후속 회의에서 결정')
card(48,210,570,180,'판단이 불가능할 때','후보 없음·센서 지연 상황에서\n하강 중지와 재탐색 검증')
card(650,210,580,180,'지형이 달라질 때','높은 착륙장과 회전한 카메라의\n고도·좌표 변환 보강')
card(48,417,570,180,'이동 중 막혔을 때','감속·정지·복구와\n실제 여유 거리·관통 여부 확인')
card(650,417,580,180,'결과를 평가할 때','착륙·대기·실패를 함께 기록\n짧은 코스부터 반복 실험')
end('현재 코드에서 후보가 없을 때 목적지로 하강하는 폴백, 고정 지면 높이와 회전 가정 같은 개선 대상을 확인했습니다. 앞으로는 이런 조건에서 안전하게 멈추고 복구하는지 검증해야 합니다. 오늘은 기술 선택을 확정하기보다는 남은 범위를 설명드리며, 구체적인 우선순위는 후속 회의에서 정하겠습니다.')

start('관람객이 상황을 바꾸는 배송 관제 체험',subtitle='시연 구상 · 대시보드는 아직 구현 전')
card(48,218,368,250,'선택','장소 카드에서 배송지 선택\n\n물류 야드 · 상점 · 광장')
card(442,218,370,250,'개입','“사람 지나가기” 버튼\n\n환경만 변화시키기')
card(838,218,392,250,'설명','카메라 위 후보와 선택점\n\n멈춤·재선택 이유 표시')
text('직접 조종보다, 드론이 스스로 판단하는 과정을 보여주는 체험',50,530,1180,29,TEAL,True)
text('안전 공간이 없으면 기다리는 것도 올바른 결과',50,588,1170,27)
end('관람객은 게임처럼 드론을 직접 조종하기보다 배송지를 선택하고 상황을 바꾸는 역할을 맡습니다. 예를 들어 사람이 지나가게 하면 드론이 카메라 변화를 보고 멈추거나 다른 착륙점을 선택하도록 하는 구상입니다. 안전한 공간이 없을 때 기다리는 것도 성공적인 안전 판단으로 보여주려 합니다. 이 화면은 향후 시연 방향이며 아직 구현된 대시보드는 아닙니다.')

start('현재 성과: 환경과 핵심 판단 프로토타입 확보',subtitle='다음 논의: 어떤 조건까지 안정적으로 증명할 것인가')
photo('metropolis/campus_aerial.jpg',48,210,625,382)
text('구축',722,213,450,28,GOLD,True);text('다양한 도시 환경과 센서 검증',722,256,495,27)
text('구현',722,328,450,28,GOLD,True);text('RGB-D 착륙점 판단과 기본 미션',722,371,495,27)
text('후속 논의',722,443,450,28,GOLD,True);text('첫 통합 코스·안전 기준·평가 범위',722,486,495,27)
text('질의응답',722,567,470,31,TEAL,True)
end('정리하면 다양한 도시 환경과 센서 검증을 마쳤고 RGB-D 착륙점 판단과 기본 미션 프로토타입을 확보했습니다. 다음 단계는 무엇을 어느 조건까지 신뢰성 있게 증명할지 범위를 정하는 일입니다. 첫 통합 코스와 안전 기준, 평가 범위에 대해 후속 논의를 진행하려 합니다. 질문 부탁드립니다.')

start('부록 · 교수님 질문에 대비한 범위 정리',kicker='APPENDIX / 발표 시간에 따라 생략')
for i,(q,a) in enumerate([('비전만 사용하는가?','장애물·착륙점 판단은 RGB-D. 자기 위치는 Gazebo odom 사용.'),('실제 드론으로 검증했는가?','현재는 위치 지정 방식의 시뮬레이션. 실기체 검증은 하지 않음.'),('움직이는 사람은 구현했는가?','하강 중 자동 진입 이벤트는 있음. 자연스러운 보행·버튼 개입은 미구현.'),('현재 92% 성공인가?','이전 정지 장면 22/24 결과. 새 도시 배송 성공률로 해석하지 않음.')]):
 y=204+i*107;text(q,49,y,1180,26,GOLD,True);text(a,49,y+43,1180,24)
end('질의응답 보조 슬라이드입니다. 위치 추정, 실제 비행, 사람 이동, 성공률의 범위를 명확히 답합니다. 검증하지 않은 기능을 구현된 것처럼 설명하지 않습니다.')

start('부록 · 근거와 재현 자료',kicker='APPENDIX / 내부 산출물 근거')
for i,(h,b) in enumerate([('현재 코드','landing_detector.py · mission_controller.py · worlds/metropolis_*.py'),('환경 검사','eval/test_metropolis.py · docs/images/metropolis/sensor_report.json'),('이전 평가','eval/results_n24/comparison.json · eval/README.md'),('영상·설계 기록','eval/record_city_tour.py · docs/PROJECT_STATUS.md · docs/HANDOFF.md')]):
 y=210+i*103;text(h,50,y,1180,26,GOLD,True);text(b,50,y+42,1180,22)
end('발표 수치와 구현 상태는 저장소 코드와 로컬 실험 원자료에 근거했습니다. 도시 캡처와 영상은 실제 Gazebo 렌더링이며, 캠퍼스의 형태는 사용자가 제공한 학교 사진을 참고했습니다. 이 발표에서는 최신 관련연구의 최초성이나 실기체 성능을 주장하지 않습니다.')

prs.save(OUT/'safe_landing_first_review.pptx')
preview=OUT/'previews';preview.mkdir(exist_ok=True)
for i,img in enumerate(slides,1):img.save(preview/f'slide_{i:02}.png')
slides[0].save(OUT/'safe_landing_first_review.pdf',save_all=True,append_images=slides[1:],resolution=96)
contact=Image.new('RGB',(1280,720),BG)
for i,img in enumerate(slides):contact.paste(img.resize((320,180)),((i%4)*320,(i//4)*180))
contact.save(OUT/'overview.jpg',quality=95)
(OUT/'speaker_notes.md').write_text('# 종합설계 1차 발표 대본\n\n본편 14장 + 부록 2장. 영상 60초 포함 약 8~10분 기준.\n\n'+'\n\n'.join(f'## {i+1:02}번 슬라이드\n\n{n}' for i,n in enumerate(notes)))
(OUT/'slide_text.json').write_text(json.dumps(alltext,ensure_ascii=False,indent=2))
print(f'Created {len(slides)} slides, editable PPTX with embedded MP4, PDF, notes and previews')
